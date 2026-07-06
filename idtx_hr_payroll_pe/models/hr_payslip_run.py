import base64
from io import BytesIO

from odoo import models, fields, _
from odoo.exceptions import UserError

try:
    import openpyxl
    from openpyxl.styles import Font
except ImportError:
    openpyxl = None


class HrPayslipRun(models.Model):
    _inherit = 'hr.payslip.run'

    # Excel de TAREO subido al crear la planilla. Al generar las boletas, se
    # cruza por DNI (identification_id) y se rellenan los "Días trabajados"
    # (worked_days) y las "Entradas salariales" (inputs) de cada boleta.
    l10n_pe_tareo_file = fields.Binary(string='Excel de Tareo', copy=False)
    l10n_pe_tareo_filename = fields.Char(string='Nombre del Tareo', copy=False)
    # Plantilla generada para descargar (pre-llenada con DNI + nombre).
    l10n_pe_tareo_template_file = fields.Binary(string='Plantilla de Tareo', copy=False)
    l10n_pe_tareo_template_filename = fields.Char(copy=False)

    # ------------------------------------------------------------------
    #  DEFINICIÓN DE COLUMNAS (única fuente de verdad: lectura + plantilla)
    #  kind: 'key' (DNI), 'info' (nombre), 'field' (campo l10n_pe_tareo_* de
    #  la boleta → worked_days/motor), 'input' (código de hr.payslip.input.type
    #  → Entradas salariales). El orden = orden de columnas del Excel.
    # ------------------------------------------------------------------
    def _pe_tareo_columns(self):
        return [
            ('DNI', 'dni', 'key'),
            ('Trabajador', 'name', 'info'),
            # Los días trabajados NO van en el tareo: premisa = todos trabajan
            # 30 días; las faltas/tardanzas se descuentan aparte. La Asistencia
            # se deriva = 30 − faltas − vacaciones − licencia (ver hr.payslip).
            # Solo se piden las horas NOCTURNAS (para la bonificación nocturna);
            # las diurnas no se usan (la Asistencia = días × jornada).
            ('Horas nocturnas', 'l10n_pe_tareo_horas_noct', 'field'),
            ('HE 25% (horas)', 'HE25_HRS', 'input'),
            ('HE 35% (horas)', 'HE35_HRS', 'input'),
            ('Faltas (días)', 'l10n_pe_tareo_faltas', 'field'),
            ('Tardanzas (horas)', 'TARDANZA_HRS', 'input'),
            # La CONDICIÓN DE TRABAJO no va en el tareo: CONDTRAB_PE = mismo
            # monto que la HE 25% (regla, result_rules['HE25_PE']).
            ('Vacaciones (días)', 'l10n_pe_tareo_vac', 'field'),
            ('Licencia c/goce (días)', 'l10n_pe_tareo_lic_cg', 'field'),
            ('Movilidad (S/)', 'MOVILIDAD', 'input'),
            ('Bonificación (S/)', 'BONIF', 'input'),
            # NO van en el tareo: el COMEDOR (DSCT_COMEDOR) lo calcula
            # idtx_hr_comedor_pe desde MySQL en compute_sheet; el ADELANTO DE
            # QUINCENA (ADEL_QUINC) es fijo para todos = parámetro
            # l10n_pe_adelanto_quincena (Ajustes ▸ Localización PE); el DOMINICAL
            # se retiró (premisa: no se paga dominical por tareo).
            ('Otros descuentos (S/)', 'OTROS_DESC', 'input'),
        ]

    # ------------------------------------------------------------------
    #  DESCARGA DE PLANTILLA (pre-llenada con los trabajadores del periodo)
    # ------------------------------------------------------------------
    def _pe_build_tareo_template(self):
        self.ensure_one()
        if openpyxl is None:
            raise UserError(_("Falta la librería 'openpyxl' en el servidor."))
        cols = self._pe_tareo_columns()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Tareo'
        ws.append([h for (h, _k, _kind) in cols])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        versions = self.env['hr.version'].browse(self._get_valid_version_ids())
        seen = set()
        for version in versions:
            emp = version.employee_id
            if not emp or emp.id in seen:
                continue
            seen.add(emp.id)
            row = [''] * len(cols)
            row[0] = emp.identification_id or ''
            row[1] = emp.name
            ws.append(row)
        bio = BytesIO()
        wb.save(bio)
        return bio.getvalue()

    def action_pe_download_tareo_template(self):
        self.ensure_one()
        data = self._pe_build_tareo_template()
        self.write({
            'l10n_pe_tareo_template_file': base64.b64encode(data),
            'l10n_pe_tareo_template_filename': 'tareo_%s.xlsx' % (self.date_start or 'plantilla'),
        })
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/hr.payslip.run/%s/l10n_pe_tareo_template_file/%s?download=true' % (
                self.id, self.l10n_pe_tareo_template_filename),
            'target': 'self',
        }

    # ------------------------------------------------------------------
    #  LECTURA DEL EXCEL → {dni: {clave: valor}}
    # ------------------------------------------------------------------
    def _pe_parse_tareo(self, b64):
        if openpyxl is None:
            raise UserError(_("Falta la librería 'openpyxl' en el servidor."))
        cols = self._pe_tareo_columns()
        try:
            wb = openpyxl.load_workbook(
                BytesIO(base64.b64decode(b64)), data_only=True, read_only=True)
        except Exception as exc:
            raise UserError(_("No se pudo leer el Excel del tareo: %s", exc))
        ws = wb.active
        rows = {}
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            if i == 0 or not r:
                continue  # cabecera / fila vacía
            if r[0] in (None, ''):
                continue
            dni = str(r[0]).strip()
            if dni.endswith('.0'):
                dni = dni[:-2]
            data = {}
            for idx, (_h, key, kind) in enumerate(cols):
                if kind in ('key', 'info'):
                    continue
                val = r[idx] if idx < len(r) else None
                try:
                    data[key] = float(val) if val not in (None, '') else 0.0
                except (TypeError, ValueError):
                    data[key] = 0.0
            rows[dni] = data
        wb.close()
        return rows

    # ------------------------------------------------------------------
    #  APLICAR TAREO A LAS BOLETAS DEL LOTE (cruce por DNI)
    # ------------------------------------------------------------------
    def _pe_apply_tareo(self):
        self.ensure_one()
        if not self.l10n_pe_tareo_file:
            return
        rows = self._pe_parse_tareo(self.l10n_pe_tareo_file)
        cols = self._pe_tareo_columns()
        input_codes = [k for (_h, k, kind) in cols if kind == 'input'] + ['HORAS_NOCT']
        type_by_code = {
            t.code: t for t in self.env['hr.payslip.input.type'].search(
                [('code', 'in', input_codes)])
        }
        applied, missing = [], []
        for slip in self.slip_ids:
            dni = (slip.employee_id.identification_id or '').strip()
            if dni.endswith('.0'):
                dni = dni[:-2]
            row = rows.get(dni)
            if not row:
                missing.append('%s (%s)' % (slip.employee_id.name, dni or 's/DNI'))
                continue
            # 1) campos del tareo (worked_days + motor)
            vals = {'l10n_pe_tareo_loaded': True}
            for (_h, key, kind) in cols:
                if kind == 'field':
                    vals[key] = row.get(key, 0.0)
            slip.write(vals)
            # 2) Días trabajados desde el tareo
            slip.write({'worked_days_line_ids': [(5, 0, 0)] + [
                (0, 0, v) for v in slip._pe_tareo_worked_day_lines()]})
            # 3) Entradas salariales: quita los inputs gestionados por tareo y recrea
            cmds = [(3, line.id) for line in slip.input_line_ids
                    if line.input_type_id.code in input_codes]
            for (_h, code, kind) in cols:
                if kind != 'input':
                    continue
                amount = row.get(code, 0.0)
                itype = type_by_code.get(code)
                if amount and itype:
                    cmds.append((0, 0, {
                        'input_type_id': itype.id, 'amount': amount, 'name': itype.name}))
            # Horas nocturnas: además del worked_day, alimenta el input HORAS_NOCT
            noct = row.get('l10n_pe_tareo_horas_noct', 0.0)
            itype_noct = type_by_code.get('HORAS_NOCT')
            if noct and itype_noct:
                cmds.append((0, 0, {
                    'input_type_id': itype_noct.id, 'amount': noct, 'name': itype_noct.name}))
            if cmds:
                slip.write({'input_line_ids': cmds})
            applied.append(slip.employee_id.name)
        # Recalcula las boletas cargadas para que queden listas.
        to_compute = self.slip_ids.filtered('l10n_pe_tareo_loaded')
        if to_compute:
            to_compute.compute_sheet()
        body = "Tareo aplicado a %d boleta(s)." % len(applied)
        if missing:
            body += "<br/>Sin fila en el tareo (no modificadas): %s" % ", ".join(missing)
        self.message_post(body=body)

    def generate_payslips(self, version_ids=None, employee_ids=None):
        res = super().generate_payslips(version_ids=version_ids, employee_ids=employee_ids)
        if self.l10n_pe_tareo_file:
            self._pe_apply_tareo()
        return res
