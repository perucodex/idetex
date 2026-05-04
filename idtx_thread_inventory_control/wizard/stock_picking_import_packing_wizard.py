import base64
import io
from collections import defaultdict

from markupsafe import Markup, escape

from odoo import _, fields, models
from odoo.exceptions import UserError


class StockPickingImportPackingWizard(models.TransientModel):
    _name = "stock.picking.import.packing.wizard"
    _description = "Importar Packing List completo en Borrador de Recepcion"

    picking_id = fields.Many2one("stock.picking", required=True, readonly=True)
    packing_file = fields.Binary(string="Packing List (Excel)", required=True)
    filename = fields.Char(string="Nombre de archivo")

    def _find_data_sheet_and_header(self, wb):
        target = {"lote", "conos", "kilos"}
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row_idx, row in enumerate(ws.iter_rows(max_row=10, values_only=True), start=1):
                cells = [str(c).strip().lower() for c in row if c is not None]
                if target.issubset(set(cells)):
                    col_map = {str(c).strip().lower(): i for i, c in enumerate(row) if c is not None}
                    return ws, row_idx, col_map
        return None, None, None

    def action_import_packing(self):
        self.ensure_one()
        picking = self.picking_id

        if not self.packing_file:
            raise UserError(_("Por favor suba un archivo de packing list."))

        try:
            import openpyxl
        except ImportError:
            raise UserError(
                _("La libreria openpyxl es requerida para importar archivos Excel. Por favor instalela.")
            )

        raw = base64.b64decode(self.packing_file)
        try:
            wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
        except Exception as exc:
            raise UserError(_("No se pudo leer el archivo Excel: %s") % str(exc))

        ws, header_row_num, col_map = self._find_data_sheet_and_header(wb)
        if not ws:
            raise UserError(
                _("No se encontro una hoja valida con columnas Lote, Conos y Kilos.")
            )

        lote_col = col_map["lote"]
        conos_col = col_map["conos"]
        kilos_col = col_map["kilos"]
        articulo_col = col_map.get("articulo")

        # raw_groups: {(product_code_upper, lot_name, cones_per_bag): [kg, ...]}
        raw_groups = defaultdict(list)

        for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if row_idx <= header_row_num:
                continue
            try:
                code_raw = str(row[articulo_col] or "").strip().upper() if articulo_col is not None else ""
                lot_raw = row[lote_col]
                cones_raw = row[conos_col]
                kilos_raw = row[kilos_col]

                if lot_raw is None or cones_raw is None or kilos_raw is None:
                    continue

                lot_name = str(lot_raw).strip()
                if not lot_name or lot_name.lower() in ("none", "total general"):
                    continue

                cones = int(float(cones_raw))
                kilos = float(kilos_raw)

                if cones <= 0 or kilos <= 0:
                    continue

                raw_groups[(code_raw, lot_name, cones)].append(kilos)
            except (ValueError, TypeError, IndexError):
                continue

        if not raw_groups:
            raise UserError(_("No se encontraron filas de datos validas en el packing list."))

        # Resolve product codes to product records (case-insensitive, is_thread only).
        all_codes = {key[0] for key in raw_groups if key[0]}
        products_by_code = {}
        for product in self.env["product.product"].search([("is_thread", "=", True)]):
            code = (product.default_code or "").strip().upper()
            if code and code in all_codes:
                products_by_code[code] = product

        missing_codes = all_codes - set(products_by_code.keys())
        found_groups = {k: v for k, v in raw_groups.items() if k[0] in products_by_code}

        if not found_groups:
            raise UserError(_(
                "No se encontraron productos en Odoo para los codigos del packing: %s"
            ) % ", ".join(sorted(all_codes)) if all_codes else _("(sin codigos detectados)"))

        # Aggregate by product: {product_code: {(lot, cones): [kg, ...]}}
        by_product = defaultdict(dict)
        for (code, lot_name, cones), weights in found_groups.items():
            by_product[code][(lot_name, cones)] = weights

        move_model = self.env["stock.move"]
        move_line_model = self.env["stock.move.line"]

        for code, lot_groups in by_product.items():
            product = products_by_code[code]
            total_qty = round(sum(sum(w) for w in lot_groups.values()), 3)

            # Find or create the stock.move for this product on the picking.
            existing_move = picking.move_ids.filtered(
                lambda m, pid=product.id: m.product_id.id == pid and m.state != "cancel"
            )[:1]

            if existing_move:
                existing_move.product_uom_qty = total_qty
                pending = existing_move.move_line_ids.filtered(
                    lambda l: l.state not in ("done", "cancel")
                )
                pending.with_context(skip_thread_combo_validation=True).unlink()
                move = existing_move
            else:
                move = move_model.create({
                    "picking_id": picking.id,
                    "product_id": product.id,
                    "product_uom_qty": total_qty,
                    "product_uom": product.uom_id.id,
                    "location_id": picking.location_id.id,
                    "location_dest_id": picking.location_dest_id.id,
                    "company_id": picking.company_id.id,
                })

            line_vals_list = []
            for (lot_name, cones_per_bag), weights in lot_groups.items():
                bag_qty = len(weights)
                total_net = round(sum(weights), 3)
                line_vals_list.append({
                    "move_id": move.id,
                    "picking_id": picking.id,
                    "product_id": product.id,
                    "location_id": picking.location_id.id,
                    "location_dest_id": picking.location_dest_id.id,
                    "lot_name": lot_name,
                    "thread_bag_qty": bag_qty,
                    "thread_cone_qty": cones_per_bag,
                    "quantity": total_net,
                    "company_id": picking.company_id.id,
                })

            move_line_model.with_context(
                skip_thread_combo_validation=True
            ).create(line_vals_list)

        if missing_codes:
            missing_codes_html = Markup("<br/>").join(
                escape(code) for code in sorted(missing_codes)
            )
            picking.message_post(
                body=Markup(
                    "<b>%s</b><br/>%s<br/>%s"
                ) % (
                    escape(_("Importacion parcial")),
                    escape(_("Productos del packing sin producto en Odoo (omitidos):")),
                    missing_codes_html,
                ),
                message_type="comment",
                subtype_xmlid="mail.mt_note",
            )

        return {"type": "ir.actions.act_window_close"}
