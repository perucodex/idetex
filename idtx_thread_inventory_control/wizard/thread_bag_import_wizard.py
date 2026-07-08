import base64
import io
from collections import defaultdict

from odoo import _, api, fields, models
from odoo.exceptions import UserError


def _clean(value):
    """Normaliza texto del Excel (quita espacios y nbsp)."""
    if value is None:
        return ""
    if isinstance(value, str):
        return value.replace("\xa0", " ").strip()
    return value


def _to_float(value):
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _to_int(value):
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return 0


class ThreadBagImport(models.Model):
    _name = "thread.bag.import"
    _description = "Importación de Bolsas de Hilo"
    _order = "create_date desc"

    name = fields.Char(string="Referencia", default=lambda s: _("Importación de bolsas"))
    import_file = fields.Binary(string="Archivo Excel", required=True)
    filename = fields.Char(string="Nombre de archivo")
    picking_type_id = fields.Many2one(
        "stock.picking.type", string="Tipo de Recepción",
        domain="[('code', '=', 'incoming')]",
        default=lambda s: s._default_picking_type(),
        help="Recepción que se creará y validará con las bolsas importadas.",
    )
    default_location_id = fields.Many2one(
        "stock.location", string="Ubicación de Ingreso (destino)", required=True,
        domain="[('usage', '=', 'internal')]",
        default=lambda s: s.env["stock.location"]._get_thread_main_location(s.env.company),
        help="Ubicación destino de las bolsas. Define el almacén/compañía y el tipo "
             "de recepción. Puedes elegir, por ejemplo, AF/Preproducción para que el "
             "hilo entre donde lo consume la orden de fabricación.",
    )
    company_id = fields.Many2one(
        "res.company", default=lambda s: s.env.company, required=True,
    )
    state = fields.Selection(
        [("draft", "Borrador"), ("done", "Importado")], default="draft",
    )
    picking_id = fields.Many2one("stock.picking", string="Recepción", readonly=True, copy=False)
    log = fields.Text(string="Resultado", readonly=True)

    @api.model
    def _default_picking_type(self):
        loc = self.env["stock.location"]._get_thread_main_location(self.env.company)
        if loc and loc.warehouse_id and loc.warehouse_id.in_type_id:
            return loc.warehouse_id.in_type_id
        return self.env["stock.picking.type"].search(
            [("code", "=", "incoming"), ("company_id", "=", self.env.company.id)], limit=1)

    @api.onchange("picking_type_id")
    def _onchange_picking_type_id(self):
        if self.picking_type_id:
            self.company_id = self.picking_type_id.company_id
            # solo sugiere ubicación si aún no eligieron una (no pisar su elección)
            if not self.default_location_id and self.picking_type_id.default_location_dest_id:
                self.default_location_id = self.picking_type_id.default_location_dest_id

    @api.onchange("default_location_id")
    def _onchange_default_location_id(self):
        """La ubicación destino manda: deriva almacén/compañía/tipo de recepción."""
        loc = self.default_location_id
        if not loc:
            return
        if loc.company_id:
            self.company_id = loc.company_id
        wh = loc.warehouse_id
        if wh and wh.in_type_id:
            self.picking_type_id = wh.in_type_id

    # ------------------------------------------------------------------
    def _find_header(self, ws):
        """Localiza la fila de encabezado y mapea columnas requeridas."""
        required = {"correl", "articulo", "lote", "kilos"}
        for row_idx, row in enumerate(ws.iter_rows(max_row=15, values_only=True), start=1):
            cells = {}
            for i, cell in enumerate(row):
                key = _clean(cell)
                if isinstance(key, str) and key:
                    cells[key.lower()] = i
            if required.issubset(set(cells.keys())):
                return row_idx, cells
        return None, None

    def _parse_rows(self, ws, header_row, cols, create_lots=True):
        """Devuelve (bag_vals_list_parcial, missing_codes, skipped_dup, skipped_invalid).

        Cada elemento de bag_vals_list es un dict con los datos crudos + product/lot
        resueltos; el move/picking se arma después. Con `create_lots=False` NO crea
        el `stock.lot` (deja `lot`=False y solo el `lot_name` de texto) — para flujos
        que difieren la creación de lotes/bolsas hasta validar la recepción.
        """
        c = cols
        Product = self.env["product.product"]
        Bag = self.env["thread.bag"]
        company = self.company_id

        product_by_code = {}
        for p in Product.search([("is_thread", "=", True)]):
            code = _clean(p.default_code)
            if code:
                product_by_code[code.upper()] = p

        existing = set(Bag.with_context(active_test=False).search(
            [("company_id", "=", company.id)]).mapped("name"))

        lot_cache = {}
        def get_lot(product, lot_name):
            key = (product.id, lot_name)
            if key not in lot_cache:
                lot_cache[key] = Bag._get_or_create_lot(product, lot_name, company)
            return lot_cache[key]

        rows = []
        missing_codes = {}
        skipped_dup = skipped_invalid = 0

        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            correl = _clean(row[c["correl"]])
            code = _clean(str(row[c["articulo"]] or "")).upper()
            lot_name = _clean(row[c["lote"]])
            net = _to_float(row[c["kilos"]])
            if not correl or correl.lower() in ("none", "total general"):
                continue
            if not code or not lot_name or net <= 0:
                skipped_invalid += 1
                continue
            if correl in existing:
                skipped_dup += 1
                continue
            product = product_by_code.get(code)
            if not product:
                missing_codes[code] = missing_codes.get(code, 0) + 1
                continue
            lot = get_lot(product, lot_name) if create_lots else False
            rows.append({
                "correl": correl, "product": product, "lot": lot,
                "lot_name": lot_name, "net": net,
                "gross": _to_float(row[c["k brutos"]]) if "k brutos" in c else 0.0,
                "tare": _to_float(row[c["tara"]]) if "tara" in c else 0.0,
                "cones": _to_int(row[c["conos"]]) if "conos" in c else 0,
                "cone_color": _clean(row[c["color cono"]]) if "color cono" in c else False,
                "packaging": _clean(row[c["tipo"]]) if "tipo" in c else False,
                "yarn_color": _clean(row[c["color"]]) if "color" in c else False,
            })
            existing.add(correl)
        return rows, missing_codes, skipped_dup, skipped_invalid

    def action_import(self):
        self.ensure_one()
        if not self.import_file:
            raise UserError(_("Suba un archivo Excel."))
        if not self.picking_type_id:
            raise UserError(_("Defina el tipo de recepción."))
        if not self.default_location_id:
            raise UserError(_("Defina la ubicación de ingreso."))
        try:
            import openpyxl
        except ImportError:
            raise UserError(_("Se requiere la librería openpyxl."))

        raw = base64.b64decode(self.import_file)
        try:
            wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
        except Exception as exc:
            raise UserError(_("No se pudo leer el Excel: %s") % exc)

        ws = header_row = cols = None
        for sheet in wb.worksheets:
            hr, cm = self._find_header(sheet)
            if hr:
                ws, header_row, cols = sheet, hr, cm
                break
        if not ws:
            raise UserError(_(
                "No se encontró una hoja con columnas Correl, Articulo, Lote y Kilos."))

        rows, missing_codes, skipped_dup, skipped_invalid = self._parse_rows(ws, header_row, cols)
        if not rows:
            raise UserError(_(
                "No hay bolsas nuevas para importar (¿duplicadas o productos inexistentes?)."))

        # La ubicación destino manda: derivamos almacén / tipo de recepción /
        # compañía / origen desde ella (aunque el tipo de recepción quede desfasado).
        dest = self.default_location_id
        wh = dest.warehouse_id
        ptype = self.picking_type_id
        if not ptype or (wh and ptype.warehouse_id and ptype.warehouse_id != wh):
            ptype = (wh.in_type_id if wh else False) or self.env["stock.picking.type"].search(
                [("code", "=", "incoming"),
                 ("company_id", "=", (dest.company_id.id or self.env.company.id))], limit=1)
        if not ptype:
            raise UserError(_("No se encontró un tipo de recepción para %s.") % dest.complete_name)
        company = dest.company_id or ptype.company_id or self.company_id
        src = ptype.default_location_src_id or self.env.ref("stock.stock_location_suppliers")

        # 1) Recepción
        picking = self.env["stock.picking"].create({
            "picking_type_id": ptype.id,
            "location_id": src.id,
            "location_dest_id": dest.id,
            "origin": self.filename or _("Importación de bolsas"),
            "company_id": company.id,
        })

        # 2) Un movimiento por producto (demanda = total neto del producto)
        Move = self.env["stock.move"]
        MoveLine = self.env["stock.move.line"]
        by_product = defaultdict(list)
        for r in rows:
            by_product[r["product"]].append(r)

        moves_by_product = {}
        for product, prod_rows in by_product.items():
            total = sum(r["net"] for r in prod_rows)
            moves_by_product[product] = Move.create({
                "picking_id": picking.id,
                "product_id": product.id,
                "product_uom": product.uom_id.id,
                "product_uom_qty": total,
                "location_id": src.id,
                "location_dest_id": dest.id,
                "company_id": company.id,
            })

        picking.action_confirm()

        # 3) Líneas de movimiento (por lote en trazados; una por producto si no)
        #    y bolsas enlazadas a la recepción.
        Bag = self.env["thread.bag"]
        bag_vals_list = []
        for product, prod_rows in by_product.items():
            move = moves_by_product[product]
            move.move_line_ids.unlink()  # quitar auto-reservas vacías
            tracked = product.tracking in ("lot", "serial")
            groups = defaultdict(list)
            for r in prod_rows:
                key = r["lot"].id if tracked else False
                groups[key].append(r)
            for key, grp in groups.items():
                MoveLine.create({
                    "move_id": move.id,
                    "product_id": product.id,
                    "product_uom_id": product.uom_id.id,
                    "lot_id": key if tracked else False,
                    "quantity": sum(g["net"] for g in grp),
                    "picked": True,
                    "location_id": src.id,
                    "location_dest_id": dest.id,
                    "company_id": company.id,
                })
            for r in prod_rows:
                bag_vals_list.append({
                    "name": r["correl"],
                    "product_id": product.id,
                    "lot_id": r["lot"].id,
                    "location_id": dest.id,
                    "cone_qty": r["cones"] or 1,
                    "gross_weight": r["gross"],
                    "tare": r["tare"],
                    "net_weight": r["net"],
                    "cone_color": r["cone_color"],
                    "packaging_type": r["packaging"],
                    "yarn_color": r["yarn_color"],
                    "company_id": company.id,
                    "receipt_picking_id": picking.id,
                    "state": "available",
                })
        Bag.create(bag_vals_list)

        # 4) Validar la recepción (genera existencia)
        validated = True
        try:
            res = picking.with_context(skip_backorder=True).button_validate()
            if isinstance(res, dict) and res.get("res_model") == "stock.backorder.confirmation":
                # No debería haber backorder (done == demanda); procesar igual.
                self.env[res["res_model"]].browse(res.get("res_id")).process()
        except Exception as exc:  # noqa: BLE001 - registrar y seguir
            validated = False
            validate_error = str(exc)

        # 5) Adjuntar el Excel a la recepción
        self.env["ir.attachment"].create({
            "name": self.filename or "import.xlsx",
            "datas": self.import_file,
            "res_model": "stock.picking",
            "res_id": picking.id,
            "type": "binary",
        })

        log_lines = [
            _("Recepción: %s") % picking.name,
            _("Bolsas creadas: %s") % len(bag_vals_list),
            _("Duplicadas (omitidas): %s") % skipped_dup,
            _("Filas inválidas (omitidas): %s") % skipped_invalid,
            _("Recepción validada (existencia generada): %s") % (_("Sí") if validated else _("NO")),
        ]
        if not validated:
            log_lines.append(_("  ⚠ Valida la recepción manualmente: %s") % validate_error)
        if missing_codes:
            log_lines.append(_("Artículos no encontrados en Odoo (omitidos):"))
            for code, n in sorted(missing_codes.items()):
                log_lines.append("   - %s (%s bolsas)" % (code, n))

        self.write({"state": "done", "picking_id": picking.id, "log": "\n".join(log_lines)})
        return {
            "type": "ir.actions.act_window",
            "res_model": "thread.bag.import",
            "res_id": self.id,
            "view_mode": "form",
            "target": "current",
        }

    def action_view_picking(self):
        self.ensure_one()
        return {
            "type": "ir.actions.act_window",
            "res_model": "stock.picking",
            "res_id": self.picking_id.id,
            "view_mode": "form",
        }
