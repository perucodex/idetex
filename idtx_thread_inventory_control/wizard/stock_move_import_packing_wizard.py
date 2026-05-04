import base64
import io
from collections import defaultdict

from odoo import _, fields, models
from odoo.exceptions import UserError


class StockMoveImportPackingWizard(models.TransientModel):
    _name = "stock.move.import.packing.wizard"
    _description = "Importar Packing List para Recepcion de Hilo"

    move_id = fields.Many2one("stock.move", string="Move", required=True, readonly=True)
    packing_file = fields.Binary(string="Packing List (Excel)", required=True)
    filename = fields.Char(string="Nombre de archivo")

    def _find_data_sheet_and_header(self, wb):
        """Return (worksheet, header_row_index_1based, col_map) for the packing sheet."""
        target = {"lote", "conos", "kilos"}
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row_idx, row in enumerate(ws.iter_rows(max_row=10, values_only=True), start=1):
                cells = [str(c).strip().lower() for c in row if c is not None]
                if target.issubset(set(cells)):
                    col_map = {}
                    for col_idx, cell in enumerate(row):
                        if cell is not None:
                            col_map[str(cell).strip().lower()] = col_idx
                    return ws, row_idx, col_map
        return None, None, None

    def action_import_packing(self):
        self.ensure_one()
        move = self.move_id

        if not self.packing_file:
            raise UserError(_("Por favor suba un archivo de packing list."))

        try:
            import openpyxl
        except ImportError:
            raise UserError(
                _("La libreria openpyxl es requerida para importar archivos Excel. Por favor instalela.")
            )

        raw = base64.b64decode(self.packing_file)
        try:
            wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
        except Exception as exc:
            raise UserError(_("No se pudo leer el archivo Excel: %s") % str(exc))

        ws, header_row_num, col_map = self._find_data_sheet_and_header(wb)
        if not ws:
            raise UserError(
                _("No se encontro una hoja valida con columnas Lote, Conos y Kilos.")
            )

        lote_col = col_map["lote"]
        conos_col = col_map["conos"]
        kilos_col = col_map["kilos"]
        articulo_col = col_map.get("articulo")

        product_code = (move.product_id.default_code or "").strip().upper()

        # Each key is (lot_name, cones_per_bag); value is list of net kg per bag.
        groups = defaultdict(list)

        for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if row_idx <= header_row_num:
                continue
            try:
                # Filter by product code when the column exists.
                if articulo_col is not None and product_code:
                    row_code = str(row[articulo_col] or "").strip().upper()
                    if row_code != product_code:
                        continue

                lot_raw = row[lote_col]
                cones_raw = row[conos_col]
                kilos_raw = row[kilos_col]

                if lot_raw is None or cones_raw is None or kilos_raw is None:
                    continue

                lot_name = str(lot_raw).strip()
                if not lot_name or lot_name.lower() in ("none", "total general"):
                    continue

                cones = int(float(cones_raw))
                kilos = float(kilos_raw)

                if cones <= 0 or kilos <= 0:
                    continue

                groups[(lot_name, cones)].append(kilos)
            except (ValueError, TypeError, IndexError):
                continue

        if not groups:
            raise UserError(_(
                "No se encontraron filas para el producto [%(code)s] en el packing list.",
                code=product_code or move.product_id.display_name,
            ))

        # Remove existing pending move lines for this move.
        pending = move.move_line_ids.filtered(lambda l: l.state not in ("done", "cancel"))
        if pending:
            pending.with_context(skip_thread_combo_validation=True).unlink()

        line_vals_list = []
        for (lot_name, cones_per_bag), weights in groups.items():
            bag_qty = len(weights)
            total_net = round(sum(weights), 3)

            line_vals_list.append(
                {
                    "move_id": move.id,
                    "picking_id": move.picking_id.id or False,
                    "product_id": move.product_id.id,
                    "location_id": move.location_id.id,
                    "location_dest_id": move.location_dest_id.id,
                    "lot_name": lot_name,
                    "thread_bag_qty": bag_qty,
                    "thread_cone_qty": cones_per_bag,
                    "quantity": total_net,
                    "company_id": move.company_id.id,
                }
            )

        # thread_cone_weight is now computed from quantity/(bags×cones) automatically.
        self.env["stock.move.line"].with_context(
            skip_thread_combo_validation=True,
        ).create(line_vals_list)

        return {"type": "ir.actions.act_window_close"}
